import os
import sys
import json
import uuid
import hashlib
from typing import Dict, Any, List

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)

from backend.app.core.config import settings
from backend.app.services.database import db
from backend.app.services.storage import storage
import fitz
from pptx import Presentation

TARGET_COURSE_ID = "a091a4aa-a333-4a0c-85ec-9af9ad8505bd"

def extract_figures_for_course(course_id: str):
    db.supabase_client = None
    db._push_to_supabase = lambda *args, **kwargs: None

    docs = db.list_all("documents", lambda d: d.get("course_id") == course_id)
    print(f"Extracting figures for course {course_id} across {len(docs)} documents...")

    # Clear existing visual elements for this course
    old_visuals = db.list_all("visual_elements", lambda v: v.get("course_id") == course_id)
    for v in old_visuals:
        db.delete("visual_elements", v["id"], save_disk=False)

    total_extracted = 0
    per_doc_counts = {}
    seen_hashes = set()

    for doc in docs:
        doc_id = doc["id"]
        fname = doc.get("filename", "")
        file_path = doc.get("storage_path")
        doc_figures = []

        if not file_path or not os.path.exists(file_path):
            print(f"Skipping {fname}: file not found on disk ({file_path})")
            per_doc_counts[fname] = 0
            continue

        if fname.lower().endswith(".pdf"):
            try:
                pdf_doc = fitz.open(file_path)
                for page_idx in range(len(pdf_doc)):
                    page_num = page_idx + 1
                    page = pdf_doc[page_idx]
                    image_list = page.get_images(full=True)
                    for img_idx, img_info in enumerate(image_list):
                        xref = img_info[0]
                        base_image = pdf_doc.extract_image(xref)
                        image_bytes = base_image.get("image")
                        image_ext = base_image.get("ext", "png")
                        w = base_image.get("width", 0)
                        h = base_image.get("height", 0)

                        if not image_bytes or len(image_bytes) < settings.FIGURE_MIN_BYTES:
                            continue
                        if w < settings.FIGURE_MIN_WIDTH or h < settings.FIGURE_MIN_HEIGHT:
                            continue

                        img_hash = hashlib.sha256(image_bytes).hexdigest()[:8]
                        if img_hash in seen_hashes:
                            continue
                        seen_hashes.add(img_hash)

                        fig_filename = f"fig_p{page_num}_{img_idx+1}_{img_hash}.{image_ext}"
                        fig_save_path = os.path.join(storage.figures_dir, fig_filename)
                        with open(fig_save_path, "wb") as f_img:
                            f_img.write(image_bytes)

                        fig_rel_path = f"/static/figures/{fig_filename}"
                        v_id = str(uuid.uuid4())
                        v_rec = {
                            "id": v_id,
                            "document_id": doc_id,
                            "course_id": course_id,
                            "element_type": "diagram",
                            "page_number": page_num,
                            "slide_number": None,
                            "caption": None,
                            "structured_description": "not analyzed",
                            "figure_path": fig_rel_path,
                            "image_url": fig_rel_path,
                            "status": "not_analyzed",
                            "created_at": "2026-10-12T12:00:00Z"
                        }
                        db.insert("visual_elements", v_rec, save_disk=False)
                        doc_figures.append(v_rec)
            except Exception as e:
                print(f"Error extracting PDF figures for {fname}: {e}")

        elif fname.lower().endswith((".pptx", ".ppt")):
            try:
                prs = Presentation(file_path)
                for slide_idx, slide in enumerate(prs.slides):
                    slide_num = slide_idx + 1
                    pic_idx = 0
                    for shape in slide.shapes:
                        if shape.shape_type == 13: # Picture
                            pic_idx += 1
                            try:
                                img = shape.image
                                image_bytes = img.blob
                                if not image_bytes or len(image_bytes) < settings.FIGURE_MIN_BYTES:
                                    continue
                                img_hash = hashlib.sha256(image_bytes).hexdigest()[:8]
                                if img_hash in seen_hashes:
                                    continue
                                seen_hashes.add(img_hash)

                                img_ext = getattr(img, "ext", "png")
                                fig_filename = f"slide_{slide_num}_{pic_idx}_{img_hash}.{img_ext}"
                                fig_save_path = os.path.join(storage.figures_dir, fig_filename)
                                with open(fig_save_path, "wb") as f_img:
                                    f_img.write(image_bytes)

                                fig_rel_path = f"/static/figures/{fig_filename}"
                                v_id = str(uuid.uuid4())
                                v_rec = {
                                    "id": v_id,
                                    "document_id": doc_id,
                                    "course_id": course_id,
                                    "element_type": "slide_visual",
                                    "page_number": None,
                                    "slide_number": slide_num,
                                    "caption": None,
                                    "structured_description": "not analyzed",
                                    "figure_path": fig_rel_path,
                                    "image_url": fig_rel_path,
                                    "status": "not_analyzed",
                                    "created_at": "2026-10-12T12:00:00Z"
                                }
                                db.insert("visual_elements", v_rec, save_disk=False)
                                doc_figures.append(v_rec)
                            except Exception:
                                pass
            except Exception as e:
                print(f"Error extracting PPTX figures for {fname}: {e}")

        per_doc_counts[fname] = len(doc_figures)
        total_extracted += len(doc_figures)
        print(f"  {fname}: {len(doc_figures)} figures extracted & saved (status='not_analyzed')")

    # Atomic database save
    db_path = settings.DB_PATH
    temp_path = db_path + f".tmp_{uuid.uuid4().hex[:8]}"
    with open(temp_path, "w", encoding="utf-8") as f:
        json.dump(db.data, f, indent=2, default=str)
    os.replace(temp_path, db_path)

    print("\n" + "=" * 60)
    print("PHASE 1.1 FIGURE EXTRACTION SUMMARY:")
    for fname, cnt in per_doc_counts.items():
        print(f"  - {fname}: {cnt} figures")
    print(f"Total visual_elements in course: {total_extracted}")
    print("All figures saved to storage/figures with status='not_analyzed'")
    print("=" * 60)

if __name__ == "__main__":
    extract_figures_for_course(TARGET_COURSE_ID)
